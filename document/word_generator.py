"""
Word 文档生成器
生成包含：患者姓名、年龄、症状、简单分析 的门诊病历
"""
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from core.exceptions import DocumentGenerateError
from utils.logger import get_logger

logger = get_logger(__name__)


class WordGenerator:
    """Word 文档生成器"""

    def __init__(self, output_dir: str = "data/output"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def generate(self, record: dict, filename: str = None) -> str:
        """生成病历 Word 文档"""
        try:
            doc = Document()

            # 设置默认字体
            style = doc.styles["Normal"]
            style.font.name = "宋体"
            style.font.size = Pt(12)

            self._add_header(doc, record)
            self._add_sections(doc, record)
            self._add_signature(doc)

            if filename is None:
                name = record.get("patient_name", "未知患者")
                date_str = datetime.now().strftime("%Y%m%d_%H%M%S")
                filename = f"病历_{name}_{date_str}.docx"

            output_path = os.path.join(self.output_dir, filename)
            doc.save(output_path)
            logger.info(f"Word 文档已生成: {output_path}")
            return output_path
        except Exception as e:
            raise DocumentGenerateError(f"Word 文档生成失败: {e}")

    def _add_header(self, doc, record: dict):
        """添加文档标题和患者基本信息"""
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("门诊病历")
        run.bold = True
        run.font.size = Pt(18)

        # 患者信息行
        info = doc.add_paragraph()
        name = record.get("patient_name") or "未填写"
        age = record.get("patient_age") or "未填写"
        date_str = datetime.now().strftime("%Y-%m-%d")
        info.add_run(f"姓名：{name}    ").bold = True
        info.add_run(f"年龄：{age}    ").bold = True
        info.add_run(f"就诊日期：{date_str}")

    def _add_sections(self, doc, record: dict):
        """按字段添加病历各章节"""
        # 症状
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run("症状：").bold = True
        symptoms = record.get("symptoms") or "未提及"
        p.add_run(str(symptoms))

        # 简单分析
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run("简单分析：").bold = True
        analysis = record.get("simple_analysis") or "未提供"
        p.add_run(str(analysis))

    def _add_signature(self, doc):
        """添加医师签名栏"""
        doc.add_paragraph()
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run("医师签名：____________")
